
from typing import Dict, List
from lib.configuration import Configuration
from baseclass.planning_parser import *
from openpyxl import load_workbook
import pandas
import re
from fnmatch import fnmatch


def nettoyer_commentaire(commentaire_brut):
    # Remplacer les espaces insécables par des espaces normaux
    commentaire = commentaire_brut.replace('\xa0', ' ')

    # Supprimer les balises automatiques et les avertissements
    if "Commentaire" in commentaire:
        # Utiliser une expression régulière pour capturer le texte après "Commentaire :"
        match = re.search(r'Commentaire\s*:\s*(.*)', commentaire, re.DOTALL)
        if match:
            commentaire = match.group(1)

    # Nettoyer les retours à la ligne, tabulations, espaces multiples
    commentaire = re.sub(r'[\r\n\t]+', ' ', commentaire)
    commentaire = re.sub(r'\s{2,}', ' ', commentaire)

    return commentaire.strip()




class AstreintePlanningParser(PlanningParser):

    SHEET_PLANNING = 'Planning'
    SHEET_ASTREINT = 'ListeAstreinte'
    SHEET_EMPLOYES = 'ListeEmployés'


    def __init__(self) -> None:
        super().__init__()

    def parse_planning(self, ignore=False) -> None:
        # lecture de tous les employes dans excel: trigramme, nom, mail, notif
        df3 = pandas.read_excel(Configuration().PATH_PLANNING_XLS, sheet_name=self.SHEET_EMPLOYES)
        print('Récup des employés...')
        employe_index: int = 0
        while employe_index < len(df3):
            employe_trigram = str(df3.iat[employe_index, 0]).upper()
            employe_name = str(df3.iat[employe_index, 1])
            if 3 < df3.shape[1]:
                notif_mail  = int(df3.iat[employe_index, 3] if str(df3.iat[employe_index, 3]) != 'nan' else 0)
                if pandas.notna(notif_mail):
                    employe_notif_mail = int(notif_mail) 
                else:
                    employe_notif_mail = 0
            else:
                employe_notif_mail = 0
            if 2 < df3.shape[1]:   
                adresse_mail = str(df3.iat[employe_index, 2] if str(df3.iat[employe_index, 2]) != 'nan' else '')
                if pandas.notna(adresse_mail) and adresse_mail != 'nan':
                    adresse_mail = adresse_mail.strip() 
                else:
                    adresse_mail = ''
            else:
                adresse_mail = ''
            if employe_trigram != None:
                employe_trigram = employe_trigram.strip()
                self.employes[employe_trigram] = EmployeInfo(employe_name, employe_notif_mail, adresse_mail )
            employe_index += 1
               
        if ignore:
            super().parse_planning()
            print('Pas d''action de parsing du planning...')
            return

        print('Parsing du planning des astreintes...')
         

        df2 = pandas.read_excel(Configuration().PATH_PLANNING_XLS, sheet_name=self.SHEET_ASTREINT)
        # Lecture de toutes les infos sur les astreintes dans excel: nom client, (primes n1 et n2), horaires
        print('Récup des infos astreintes...')
        line_index: int = 0
        column_name: int = 0
        column_n1: int = 3
        column_n2: int = 4
        column_horaires: int = 7
        column_packAuto: int = 8
        column_packInfo: int = 9
        company_name = str(df2.iat[line_index, column_name])
        while company_name != "nan":
            prime_n1 = int(df2.iat[line_index, column_n1] if str(df2.iat[line_index, column_n1]) != 'nan' else 0)
            prime_n2 = int(df2.iat[line_index, column_n2] if str(df2.iat[line_index, column_n2]) != 'nan' else 0)
            horaires = str(df2.iat[line_index, column_horaires] if str(df2.iat[line_index, column_horaires]) != 'nan' else '')
            nom_PackAuto = str(df2.iat[line_index, column_packAuto] if str(df2.iat[line_index, column_packAuto]) != 'nan' else '')
            nom_PackInfo = str(df2.iat[line_index, column_packInfo] if str(df2.iat[line_index, column_packInfo]) != 'nan' else '')
            self.primes_astreinte[company_name] = PrimeAstreinte(company_name, prime_n1, prime_n2, horaires, nom_PackAuto, nom_PackInfo)
            line_index += 1
            company_name = str(df2.iat[line_index, column_name])

            
            
        df = pandas.read_excel(Configuration().PATH_PLANNING_XLS, sheet_name=self.SHEET_PLANNING)
        print('Récup des astreintes...')
        wb = load_workbook(Configuration().PATH_PLANNING_XLS, data_only=True)
        ws = wb[self.SHEET_PLANNING]
        print(' et des commentaires...')

        commentaires = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    # Initialiser le sous-dictionnaire si la clé n'existe pas encore
                    if cell.row not in commentaires:
                        commentaires[cell.row] = {}
                    commentaires[cell.row][cell.column] = nettoyer_commentaire(cell.comment.text)
        wb.close()

        # memorisation du type d'astreinte pour chaque client (N1, N1+N2) et des numéros de colonne contenant l'info
        # sera utilisé plus tard
        astreinte_par_colonne = {}
        pack_par_colonne = {}
        c=2
    
        for i in range(1,150):
            nme = str(df.iat[0, c])  # nom du client
            n1 = str(df.iat[2, c])  # N1 normalement
            n2 = str(df.iat[2, c+1]) # N2 ou N1 si 1 seul N1 pour le client précédent

            exp1 = str(df.iat[3, c])    # Exp Auto Normalement ou N/A si pas de pack expert
            exp2 = str(df.iat[3, c+1]) # Exp Info Normalement ou Exp Auto ou N/A

            if nme == 'nan':
                break

            if nme != None:
                nme = nme.strip()
            if n1 != None:
                n1 = n1.strip()
            if n2 != None:
                n2 = n2.strip()
            if exp1 != None:
                exp1 = exp1.strip()
            if exp2 != None:
                exp2 = exp2.strip()


            if nme == None and n2 == None:
                break

            if nme == None:
                continue

            found = True
            lvls = []
            pack = []
            if n1 == 'N3':
                lvls = ['N3']
            elif n1 == 'N1' and n2=='N2':
                lvls = ['N1','N2']
                if exp1 == 'Exp Auto' and exp2 == 'Exp Info':
                    pack= ['Auto','Info']
                elif exp1 == 'Exp Info' and exp2 == 'Exp Auto':
                    pack= ['Info','Auto']
            elif n2 == 'N1':            
                lvls = ['N1_only']
                if exp1 == 'Exp Auto' :
                    pack= ['Auto']
                elif exp1 == 'Exp Info':
                    pack= ['Info']
            else:
                found = False
                print('ERREUR: Detection site N1 N2')

            if found == True:                           
                obj={
                    'col':c,
                    'name':nme,
                    'levels':lvls,
                    'packExpert':pack
                    }
            
                astreinte_par_colonne[c]=obj
            
                c+=len(lvls)

        # mémorisation du numéro de ligne excel contenant les infos de chaque numéro de semaine
        # sera utilisé plus tard
        print('Récup des semaines d''astreinte...')
        pattern = re.compile(r'^S([0-9]+)$')
        pattern_pack = re.compile(r'^S([0-9]+)\sExp\.$')
        numero_ligne_par_numero_semaine = {}
        numero_ligne_par_numero_semaine_pack = {}
        for row_index in range(5,150):
            txt = str(df.iat[row_index, 1])
            match = pattern.match(str(txt))
            if match != None:
            
                numero_semaine = int(match.group(1))
                if numero_semaine in numero_ligne_par_numero_semaine:
                    print('[WARNING] - Doublon de semaine %d dans le fichier source' % (numero_semaine))
                
                numero_ligne_par_numero_semaine[numero_semaine]=row_index
            else:
                match = pattern_pack.match(str(txt))
                if match != None:
            
                    numero_semaine = int(match.group(1))
                    if numero_semaine in numero_ligne_par_numero_semaine_pack:
                        print('[WARNING] - Doublon de semaine PackExpert %d dans le fichier source' % (numero_semaine))
                
                    numero_ligne_par_numero_semaine_pack[numero_semaine]=row_index


        # Parcours du planning en utilisant les infos memorisees precedemment (astreinte_par_colonne et numero_ligne_par_numero_semaine)
        # Objectif: générer un dictionnaire Trigramme -> Num Semaine -> Liste Astreinte
        print('Récup du planning...')
        self.affectation_astreintes: Dict[str, Dict[int, List[AstreinteInfo]]] = dict()
        for week_nbr,row_index in numero_ligne_par_numero_semaine.items():
            lst = []
            row_index_pack = numero_ligne_par_numero_semaine_pack[week_nbr] #rang excel des infos pack (juste en dessous normalement)
            for c,ast in astreinte_par_colonne.items():
                inx=0
                indx_binome=1
                for lvl in ast['levels']:
                    valeurExcel = df.iat[row_index, c+inx] #Nom du N1 ou du N2
                    # Bypass des astreintes N/A
                    if pandas.isna(valeurExcel):
                        continue
                    trigramme=str(valeurExcel).strip().upper()
                    
                    # Gestion des binômes
                    if len(ast['levels']) > 1 :
                        binome=str(df.iat[row_index, c+indx_binome-inx]).upper()
                    else:
                        binome = ''
                    # Gestion des commentaires
                    offsetLigne = 2
                    offsetColonne = 1
                    commentaire = commentaires.get(row_index+offsetLigne, {}).get(c+offsetColonne+inx, '')
                    # Gestion des trigrammes multiples
                    list_trigrams = trigramme.split('/')
                    current_lvl = lvl
                    for un_trigram in list_trigrams:
                        item = AstreinteInfo(ast['name'], current_lvl, binome, commentaire, week_nbr)
                        if un_trigram not in self.affectation_astreintes.keys():
                            self.affectation_astreintes[un_trigram] = dict()

                        if week_nbr not in self.affectation_astreintes[un_trigram].keys():
                            self.affectation_astreintes[un_trigram][week_nbr] = list()

                        self.affectation_astreintes[un_trigram][week_nbr].append(item)
                        # Pour les autres trigrammes, on est en exception du niveau
                        if(not current_lvl.startswith('E_')):
                            current_lvl = 'E_' + lvl

                    # Ajout des packs expert 
                    # si existant
                    if len(ast['packExpert']) > inx:
                        currentPack = ast['packExpert'][inx]
                        valeurPack_Excel = df.iat[row_index_pack, c+inx] #Trigramme du premier pack expert
                        if not pandas.isna(valeurPack_Excel):
                            trigrammeExp=str(valeurPack_Excel).strip().upper()
                            itemPack = AstreinteInfo(ast['name'], currentPack, '', '', week_nbr)

                            if trigrammeExp not in self.affectation_astreintes.keys():
                                self.affectation_astreintes[trigrammeExp] = dict()

                            if week_nbr not in self.affectation_astreintes[trigrammeExp].keys():
                                self.affectation_astreintes[trigrammeExp][week_nbr] = list()

                            self.affectation_astreintes[trigrammeExp][week_nbr].append(itemPack)

                    inx+=1
                    
                    



        # Retravail du dictionnaire pour ajouter les astreintes affectées à plusieurs personnes
        for trigram, ast_by_week in self.affectation_astreintes.items():
            if len(trigram) <= 6:
                continue
            #contained_trigrams = [t for t in self.affectation_astreintes if t in trigram and t != trigram]
            contained_trigrams = trigram.split('/')
            for linked_trigram in contained_trigrams:
                for week_number, list_astreints in ast_by_week.items():
                    if week_number not in self.affectation_astreintes[linked_trigram].keys():
                        self.affectation_astreintes[linked_trigram][week_number] = list()

                    for astreint in list_astreints:
                        if astreint not in self.affectation_astreintes[linked_trigram][week_number]:
                            # Ajouter l'astreinte à la liste de l'autre trigramme
                            self.affectation_astreintes[linked_trigram][week_number].append(astreint)

        # Appel à la méthode de base en fin d'import pour mettre à jour les flags internes
        super().parse_planning()

    def check_attendee_constraints(self, ctrl_week_number: int) -> List[str]:
        inconsistencies = list()
        user = Configuration().attendee_trigram
        for constraint in Configuration().attendee_constraints:
            count = 0
            for week_number, astreints in self.get_astreintes(user).items():
                if week_number != ctrl_week_number:
                    continue

                for astreint in astreints:
                    if fnmatch(astreint.company, constraint.company):
                        count += 1
                        if count > constraint.max:
                            print(f"Contrainte non respectée pour la semaine {week_number} : {constraint}")
                            inconsistencies.append(f"Contrainte non respectée pour la semaine {week_number} : {constraint}")
                            break

        # Astreintes limitées à ce qui est listé dans les contraintes
        for week_number, astreints in self.get_astreintes(user).items():
            if week_number != ctrl_week_number:
                continue

            not_wanted = [astr for astr in astreints]
            for astreint in astreints:
                for constraint in Configuration().attendee_constraints:
                    if constraint.company == "*":
                        if len(Configuration().attendee_constraints) == 1:
                            not_wanted.pop(not_wanted.index(astreint))
                        continue
                    if fnmatch(astreint.company, constraint.company):
                        not_wanted.pop(not_wanted.index(astreint))

            if not_wanted:
                print(f"Contrainte non respectée pour la semaine {week_number} : Entreprise(s) non souhaitée(s) {[item.company for item in not_wanted]}")
                inconsistencies.append(f"Contrainte non respectée pour la semaine {week_number} : Entreprise(s) non souhaitée(s) {[item.company for item in not_wanted]}")
                break

        return inconsistencies